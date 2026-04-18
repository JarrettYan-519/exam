"""Extract plain text from uploaded resume files."""

from __future__ import annotations

import io
import os

from fastapi import HTTPException, status

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def extract_text(filename: str, data: bytes, max_mb: int) -> str:
    if len(data) > max_mb * 1024 * 1024:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"文件过大，最大支持 {max_mb} MB",
        )

    ext = os.path.splitext(filename.lower())[1]
    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"不支持的文件类型 {ext}。仅支持 {sorted(SUPPORTED_EXTENSIONS)}",
        )

    try:
        if ext in {".txt", ".md"}:
            return data.decode("utf-8", errors="replace").strip()
        if ext == ".pdf":
            return _extract_pdf(data)
        if ext == ".docx":
            return _extract_docx(data)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"解析文件失败：{exc}",
        )

    # unreachable
    raise HTTPException(status_code=500, detail="unreachable")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    result = "\n\n".join(parts).strip()
    if not result:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="PDF 未能抽取到文本（可能是扫描件，请改为文本/Markdown 粘贴）",
        )
    return result


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    # also pull table cells
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs).strip()
